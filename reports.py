"""Business report/export helpers for INSITEVA."""
from pathlib import Path
import json
import pandas as pd

def report_payload(dataset_name, df, question="", answer="", evidence=None, explanation=None, recommendations=None):
    return {
        "dataset": dataset_name or "Dataset",
        "generated_at": pd.Timestamp.now().isoformat(),
        "rows": int(len(df)) if df is not None else 0,
        "columns": int(len(df.columns)) if df is not None else 0,
        "question": question or "",
        "answer": answer or "",
        "explanation": explanation or "",
        "recommendations": recommendations or [],
        "evidence": evidence or {},
    }

def export_excel(path, dataset_name, df, payload):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Data")
        summary = pd.DataFrame([
            ["Dataset", dataset_name], ["Rows", payload["rows"]],
            ["Columns", payload["columns"]], ["Question", payload["question"]],
            ["Answer", payload["answer"]], ["Explanation", payload["explanation"]],
        ], columns=["Field","Value"])
        summary.to_excel(writer, index=False, sheet_name="AI Summary")
        ev = payload.get("evidence") or {}
        pd.DataFrame({"Evidence": [json.dumps(ev, indent=2, default=str)]}).to_excel(writer,index=False,sheet_name="Evidence")

def export_html(path, payload, df):
    esc=lambda x: str(x).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")
    preview=df.head(100).to_html(index=False, classes="data") if df is not None else ""
    html=f"""<!doctype html><html><head><meta charset="utf-8"><title>INSITEVA Report</title>
<style>body{{font-family:Arial,sans-serif;margin:40px;color:#17202a}}h1{{margin-bottom:4px}}.meta{{color:#607080}}.card{{padding:16px;margin:14px 0;border:1px solid #d8e0e8;border-radius:10px}}.data{{border-collapse:collapse;width:100%;font-size:12px}}.data th,.data td{{border:1px solid #ddd;padding:6px}}.data th{{background:#eef4f8}}</style></head>
<body><h1>INSITEVA — Analyst Report</h1><div class="meta">{esc(payload["dataset"])} • {payload["generated_at"]}</div>
<div class="card"><b>Question</b><p>{esc(payload["question"])}</p><b>Answer</b><p>{esc(payload["answer"])}</p></div>
<div class="card"><b>AI Explanation</b><p>{esc(payload["explanation"])}</p></div>
<div class="card"><b>Dataset</b><p>{payload["rows"]:,} rows × {payload["columns"]} columns</p></div>
<div class="card"><b>Recommendations</b><ul>{''.join("<li>"+esc(x)+"</li>" for x in payload.get("recommendations",[])) or "<li>No recommendations recorded.</li>"}</ul></div>
<div class="card"><b>Evidence</b><pre>{esc(json.dumps(payload.get("evidence",{}),indent=2,default=str))}</pre></div>
<h2>Data Preview</h2>{preview}</body></html>"""
    Path(path).write_text(html,encoding="utf-8")

def export_pdf(path,payload,df):
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    styles=getSampleStyleSheet()
    doc=SimpleDocTemplate(path,pagesize=A4,rightMargin=36,leftMargin=36,topMargin=36,bottomMargin=36)
    story=[Paragraph("INSITEVA — Analyst Report",styles["Title"]),Paragraph(f'{payload["dataset"]} • {payload["generated_at"]}',styles["Normal"]),Spacer(1,12)]
    for title,key in [("Question","question"),("Answer","answer"),("AI Explanation","explanation")]:
        story += [Paragraph(title,styles["Heading2"]),Paragraph(str(payload[key]).replace("&","&amp;"),styles["BodyText"]),Spacer(1,8)]
    story += [Paragraph("Dataset",styles["Heading2"]),Paragraph(f'{payload["rows"]:,} rows × {payload["columns"]} columns',styles["BodyText"]),Spacer(1,8)]
    story += [Paragraph("Recommendations",styles["Heading2"])]
    for r in payload.get("recommendations",[]): story.append(Paragraph("• "+str(r),styles["BodyText"]))
    story += [Spacer(1,10),Paragraph("Evidence",styles["Heading2"]),Paragraph("<font size='7'>"+str(json.dumps(payload.get("evidence",{}),default=str))[:6000].replace("&","&amp;").replace("<","&lt;")+"</font>",styles["BodyText"])]
    if df is not None and len(df):
        sample=df.head(20).copy().astype(str); data=[list(sample.columns)]+sample.values.tolist()
        t=Table(data,repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e8eef3")),("GRID",(0,0),(-1,-1),.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),6)]))
        story += [PageBreak(),Paragraph("Data Preview",styles["Heading2"]),t]
    doc.build(story)

def export_pptx(path,payload):
    from pptx import Presentation
    prs=Presentation()
    def as_text(value):
        if isinstance(value,(dict,list,tuple)):
            return json.dumps(value,indent=2,default=str)
        return str(value)
    def slide(title,body):
        s=prs.slides.add_slide(prs.slide_layouts[1]);s.shapes.title.text=title;s.placeholders[1].text=as_text(body)
    slide("INSITEVA — Analyst Report",f'{payload["dataset"]}\n{payload["rows"]:,} rows × {payload["columns"]} columns')
    slide("Question & Answer",f'{payload["question"]}\n\n{payload["answer"]}')
    slide("AI Explanation",payload["explanation"] or "No explanation recorded.")
    rec="\n".join("• "+str(x) for x in payload.get("recommendations",[])) or "No recommendations recorded."
    slide("Recommendations",rec)
    slide("Evidence",json.dumps(payload.get("evidence",{}),indent=2,default=str)[:5000])
    prs.save(path)
